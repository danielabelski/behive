"""Review export row fidelity and formula-injection refusal."""
import importlib
from pathlib import Path
import sys
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))


def test_export_sheet_rows_and_untrusted_cells(tmp_path):
    assert (ROOT/'tools/eval/export_xlsx.py').exists(), 'review exporter missing'
    m=importlib.import_module('tools.eval.export_xlsx')
    rows=[{'url':'https://a','title':'=1+1','content':{'text':'x'*900},'sol':2,'claude':1,'opus':None,'final':1.5,'disputed':False,'reasons':{'sol':'useful'}}]
    p=m.export_topics({'T2':{'items':rows}},tmp_path/'gold.xlsx')
    import openpyxl
    w=openpyxl.load_workbook(p)
    assert w.sheetnames==['T2']
    s=w['T2']
    assert s.max_row==2
    assert s['A2'].value=='https://a'
    assert s['B2'].data_type=='s'
    assert s['B2'].value=='=1+1'
    assert len(s['C2'].value)==500
    assert s['G2'].value==1.5
