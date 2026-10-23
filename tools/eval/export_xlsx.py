"""One topic per review sheet; CSV fallback when openpyxl is unavailable."""
from __future__ import annotations

import argparse
import csv
import importlib
import json
from pathlib import Path
import re
import sys

COLUMNS = ['url', 'title', 'excerpt', 'sol', 'claude', 'opus', 'final', 'disputed', 'sol_reason', 'claude_reason', 'opus_reason', 'evidence_type', 'fetch_status']


def _rows(gold):
    for row in gold['items']:
        reasons = row.get('reasons', {})
        content = row.get('content', {})
        yield [row.get('url', ''), row.get('title', ''), (content.get('text') or row.get('snippet') or row.get('claim') or '')[:500],
               row.get('sol'), row.get('claude'), row.get('opus'), row.get('final'), row.get('disputed'),
               reasons.get('sol'), reasons.get('claude'), reasons.get('opus'), row.get('evidence_type', ''), content.get('status', '')]


def _openpyxl():
    try:
        return importlib.import_module('openpyxl')
    except ImportError:
        # openpyxl is pure Python: reuse the already installed Hermes copy; no installs.
        for directory in ('/usr/local/lib/hermes-agent/venv/lib/python3.11/site-packages',):
            if Path(directory, 'openpyxl').is_dir() and directory not in sys.path:
                sys.path.append(directory)
        try:
            return importlib.import_module('openpyxl')
        except ImportError:
            return None


def export_topics(topics, output):
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    openpyxl = _openpyxl()
    if openpyxl is None:
        directory = output.with_suffix('')
        directory.mkdir(parents=True, exist_ok=True)
        for name, gold in topics.items():
            safe = re.sub(r'[^\w.-]', '_', name)
            with (directory / (safe + '.csv')).open('w', newline='', encoding='utf-8-sig') as stream:
                writer = csv.writer(stream)
                writer.writerow(COLUMNS)
                for row in _rows(gold):
                    # CSV cannot type cells; neutralize formula execution explicitly.
                    writer.writerow(["'" + v if isinstance(v, str) and v.startswith(('=', '+', '-', '@')) else v for v in row])
        return directory
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name, gold in topics.items():
        sheet = workbook.create_sheet(re.sub(r'[\[\]:*?/\\]', '_', name)[:31])
        sheet.append(COLUMNS)
        for values in _rows(gold):
            sheet.append([ILLEGAL_CHARACTERS_RE.sub('', v) if isinstance(v, str) else v for v in values])
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = 's'  # untrusted pages must never become Excel formulas
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='244354')
        widths = [60, 48, 90, 9, 9, 9, 9, 12, 45, 45, 45, 20, 24]
        for i, width in enumerate(widths, 1):
            sheet.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width
        sheet.freeze_panes = 'D2'
        sheet.auto_filter.ref = sheet.dimensions
    if not topics:
        workbook.create_sheet('No labels')
    workbook.save(output)
    return output


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('gold_dir')
    p.add_argument('out')
    a = p.parse_args()
    topics = {p.stem: json.loads(p.read_text()) for p in sorted(Path(a.gold_dir).glob('T*.json'))}
    print(export_topics(topics, a.out))


if __name__ == '__main__':
    main()
